import random
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
import os 

os.chdir('/home/saul/Business')
workbook = Workbook()
sheet = workbook.active

# Let's create some sample sales data
rows = [
    ["", "January", "February", "March", "April",
    "May", "June", "July", "August", "September",
     "October", "November", "December"],
    [1, ],
    [2, ],
    [3, ],
]

for row in rows:
    sheet.append(row)
 
for row in sheet.iter_rows(min_row=2,
                           max_row=4,
                           min_col=2,
                           max_col=13):
    print("Row ", row)
    for cell in row:
        cell.value = random.randrange(5, 100)
        
chart = LineChart()
data = Reference(worksheet=sheet,
                 min_row=2,
                 max_row=4,
                 min_col=1,
                 max_col=13)
chart.add_data(data, from_rows=True, titles_from_data=True)
sheet.add_chart(chart, "C6")
workbook.save("line_chart.xlsx")
